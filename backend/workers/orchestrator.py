"""
Orchestrateur pour la génération de rapports.
Coordonne les 3 étapes : extraction, génération, rendu.
"""
import json
import logging
import tempfile
import shutil
from collections import Counter
from pathlib import Path
from typing import Dict, Any, Optional, Callable, List
from dataclasses import dataclass

# Import des modules CLI du dossier CLIENTS et core
import sys
project_root = Path(__file__).parent.parent.parent
sys.path.insert(0, str(project_root / "CLIENTS"))
sys.path.insert(0, str(project_root))

from extract_sources import extract_one, walk_files, ExtractedDoc
from core.generate import generate_fields
from core.template_fields import build_field_specs, extract_placeholders_from_docx
from core.location_date import build_location_date
from core.avs import detect_avs_number
from core.render import render_report

logger = logging.getLogger(__name__)


@dataclass
class ReportGenerationParams:
    """Paramètres pour la génération de rapport."""
    # Chemins
    client_dir: Path
    template_path: Path
    output_path: Path
    
    # Identité
    name: str = ""
    surname: str = ""
    civility: str = "Monsieur"
    avs_number: str = ""
    
    # Lieu et date
    location_city: str = ""
    location_date: str = ""
    auto_location_date: bool = True
    date_format: str = "%d/%m/%Y"
    
    # LLM
    llm_host: str = "http://localhost:11434"
    llm_model: str = "mistral:latest"
    temperature: float = 0.2
    topk: int = 10
    top_p: float = 0.9
    
    # Extraction
    extract_method: str = "auto"
    source_file: Optional[str] = None
    enable_soffice: bool = False
    
    # Filtres
    include_filters: list = None
    exclude_filters: list = None
    
    # Options
    force_reextract: bool = False
    export_pdf: bool = False

    # Audio RAG
    # Si True, transcrit automatiquement (Whisper local) les audios du client
    # qui n'ont pas encore de manifest ingested_audio/*.json.
    # Objectif: éviter de devoir lancer manuellement "ingest-local" avant chaque rapport.
    auto_ingest_audio: bool = True
    max_audio_ingest_files: int = 25


class ReportOrchestrator:
    """Orchestrateur pour la génération de rapports."""
    
    def __init__(self, params: ReportGenerationParams, progress_callback: Optional[Callable] = None):
        self.params = params
        self.progress_callback = progress_callback
        self.temp_dir = None
        self.logs = []  # Historique des logs
        self.field_progress: Dict[str, Dict[str, Any]] = {}
        self.field_order: List[str] = []
        self.field_progress_version: int = 0
        # Statistiques sources (utile pour debug RAG / extraction)
        self.source_stats: Optional[Dict[str, Any]] = None
        
    def _log_progress(self, status: str, message: str, progress: Optional[float] = None, *, include_fields: bool = False):
        """Log et appel du callback de progression."""
        from datetime import datetime
        logger.info(f"[{status}] {message}")
        
        # Ajouter aux logs
        log_entry = {
            "phase": status,
            "message": message,
            "progress": progress,
            "timestamp": datetime.now().isoformat()
        }
        self.logs.append(log_entry)
        
        # Callback avec logs
        if self.progress_callback:
            payload: Dict[str, Any] = {
                "status": status,
                "message": message,
                "progress": progress,
                "logs": self.logs,
                "timestamp": log_entry["timestamp"]
            }
            # Toujours inclure les stats de sources si disponibles (frontend peut les afficher).
            if self.source_stats:
                payload["source_stats"] = self.source_stats
            if include_fields:
                payload.update(
                    {
                        "field_progress": self.field_progress,
                        "field_order": self.field_order,
                        "field_progress_version": self.field_progress_version,
                    }
                )
            self.progress_callback(payload)

    def _init_field_progress(self, fields: List[Dict[str, Any]]) -> None:
        """Initialise la structure de progression par champ (style Streamlit)."""
        from datetime import datetime

        self.field_progress = {}
        self.field_order = []
        now_iso = datetime.now().isoformat()
        for field in fields:
            key = field.get("key")
            if not key:
                continue
            label = field.get("label") or field.get("title") or str(key).replace("_", " ").title()
            self.field_order.append(key)
            self.field_progress[key] = {
                "label": label,
                "stage": "pending",
                "message": "En attente",
                "updated_at": now_iso,
            }
        self.field_progress_version += 1

    def _update_field_progress(self, field_key: str, stage: str, message: str) -> None:
        """Met à jour la progression d'un champ + pousse l'update au callback."""
        from datetime import datetime

        if field_key not in self.field_progress:
            self.field_progress[field_key] = {
                "label": str(field_key).replace("_", " ").title(),
                "stage": "pending",
                "message": "",
                "updated_at": datetime.now().isoformat(),
            }
            self.field_order.append(field_key)

        self.field_progress[field_key]["stage"] = stage
        self.field_progress[field_key]["message"] = message
        self.field_progress[field_key]["updated_at"] = datetime.now().isoformat()
        self.field_progress_version += 1

        total = max(len(self.field_order), 1)
        done = 0
        for key in self.field_order:
            stg = (self.field_progress.get(key) or {}).get("stage")
            if stg in {"done", "warning", "error"}:
                done += 1

        # Progression globale pendant GENERATING : 0.4 → 0.7
        percent = 0.4 + 0.3 * (done / total)
        percent = min(max(percent, 0.4), 0.7)

        if self.progress_callback:
            self.progress_callback(
                {
                    "status": "GENERATING",
                    "message": f"LLM [{field_key}] {stage}: {message}",
                    "progress": percent,
                    "field_progress": self.field_progress,
                    "field_order": self.field_order,
                    "field_progress_version": self.field_progress_version,
                    "timestamp": datetime.now().isoformat(),
                }
            )
    
    def run(self) -> Dict[str, Any]:
        """
        Exécute le pipeline complet de génération de rapport.
        
        Returns:
            dict: Résultat avec status, output_path, etc.
        """
        try:
            # Création du répertoire temporaire
            self.temp_dir = Path(tempfile.mkdtemp(prefix="rapport_"))
            
            # ÉTAPE 1: Extraction
            self._log_progress("EXTRACTING", "Extraction des documents sources...", 0.1)
            extracted_path = self._extract_sources()
            self._log_progress("EXTRACTING", f"{extracted_path.stat().st_size} octets extraits", 0.3)
            
            # ÉTAPE 2: Génération des champs via LLM
            self._log_progress("GENERATING", "Génération des champs avec le LLM...", 0.4)
            answers_path = self._generate_fields(extracted_path)
            self._log_progress("GENERATING", "Champs générés avec succès", 0.7)
            
            # ÉTAPE 3: Rendu du document Word
            self._log_progress("RENDERING", "Création du document Word...", 0.8)
            self._render_document(answers_path)
            self._log_progress("RENDERING", "Document Word créé", 0.9)
            
            # ÉTAPE 4 (optionnelle): Export PDF
            if self.params.export_pdf:
                self._log_progress("EXPORTING", "Export PDF...", 0.95)
                self._export_pdf()
                self._log_progress("EXPORTING", "PDF exporté", 1.0)
            
            self._log_progress("COMPLETED", f"Rapport généré: {self.params.output_path}", 1.0)
            
            return {
                "status": "success",
                "output_path": str(self.params.output_path),
                "client_name": self.params.client_dir.name,
                "temp_files": {
                    "extracted": str(extracted_path),
                    "answers": str(answers_path)
                }
            }
            
        except Exception as e:
            logger.exception("Erreur lors de la génération du rapport")
            self._log_progress("FAILED", f"Erreur: {str(e)}", None)
            return {
                "status": "failed",
                "error": str(e),
                "error_type": type(e).__name__
            }
    
    def _extract_sources(self) -> Path:
        """Extraction des documents sources."""
        extracted_path = self.temp_dir / "extracted.json"

        def _audio_deps_ok() -> bool:
            # Dépendances nécessaires à faster-whisper (ffmpeg/ffprobe + module Python)
            if not shutil.which("ffmpeg") or not shutil.which("ffprobe"):
                return False
            try:
                import faster_whisper  # type: ignore  # noqa: F401
            except Exception:
                return False
            return True

        def _maybe_auto_ingest_audio() -> Dict[str, Any]:
            """Transcrit automatiquement les audios manquants (idempotent via manifests JSON).

            Retourne un petit résumé pour logs / debug.
            Ne lève pas en cas d'échec: l'extraction de documents texte doit continuer.
            """

            summary: Dict[str, Any] = {
                "enabled": bool(getattr(self.params, "auto_ingest_audio", False)),
                "queued": 0,
                "processed": 0,
                "skipped": 0,
                "errors": 0,
            }

            if not getattr(self.params, "auto_ingest_audio", False):
                return summary

            # Si deps manquantes, on ne bloque pas le rapport.
            if not _audio_deps_ok():
                summary["skipped_reason"] = "missing_deps"
                return summary

            client_dir = self.params.client_dir
            ingested_dir = client_dir / "sources" / "ingested_audio"
            ingested_dir.mkdir(parents=True, exist_ok=True)

            # Détecter les audios déjà ingérés (via manifests JSON)
            seen_audio_paths: set[str] = set()
            try:
                for mf in ingested_dir.glob("*.json"):
                    try:
                        payload = json.loads(mf.read_text(encoding="utf-8"))
                        ap = payload.get("audio_path")
                        if isinstance(ap, str) and ap:
                            seen_audio_paths.add(str(Path(ap).expanduser().resolve()))
                    except Exception:
                        continue
            except Exception:
                # Non bloquant
                pass

            # Scanner les audios dans tout le dossier client (hors ingested_audio)
            allowed_exts = {".m4a", ".mp3", ".wav"}
            audio_files: list[Path] = []
            try:
                for p in client_dir.rglob("*"):
                    if not p.is_file():
                        continue
                    if p.suffix.lower() not in allowed_exts:
                        continue
                    rp = p.expanduser().resolve()
                    # Éviter de ré-ingérer ce qu'on a généré
                    if "sources/ingested_audio" in str(rp).replace("\\", "/"):
                        continue
                    if str(rp) in seen_audio_paths:
                        summary["skipped"] += 1
                        continue
                    audio_files.append(rp)
            except Exception:
                # Non bloquant
                return summary

            if not audio_files:
                return summary

            max_files = int(getattr(self.params, "max_audio_ingest_files", 25) or 0)
            if max_files > 0:
                audio_files = audio_files[:max_files]

            # Import local (pour éviter d'alourdir le démarrage worker + facilités de tests)
            try:
                from script_ai.rag.ingest_audio import ingest_audio_file  # type: ignore
            except Exception:
                summary["skipped_reason"] = "import_error"
                return summary

            self._log_progress(
                "EXTRACTING",
                f"🎙️ Audio: transcription automatique de {len(audio_files)} fichier(s) (max={max_files})…",
                0.11,
            )

            for p in audio_files:
                try:
                    ingest_audio_file(
                        str(p),
                        source_id=client_dir.name,
                        extra_metadata={"origin": "auto_report", "relative_path": str(p.relative_to(client_dir)) if str(p).startswith(str(client_dir)) else str(p)},
                    )
                    summary["processed"] += 1
                except Exception:
                    summary["errors"] += 1
                    # On continue, le rapport ne doit pas échouer pour l'audio.
                    continue

            return summary

        def _attempt_extract(input_dir: Path) -> list[ExtractedDoc]:
            # Liste des fichiers
            files = walk_files(input_dir)
            if not files:
                raise ValueError(f"Aucun fichier trouvé dans {input_dir}")

            # Stats de sources (debug RAG: vérifier qu'on scanne le bon dossier + que l'audio ingéré est présent)
            ext_counts = Counter()
            for fp in files:
                ext = fp.suffix.lower() if fp.suffix else "(sans_ext)"
                ext_counts[ext] += 1

            ingested_dir = self.params.client_dir / "sources" / "ingested_audio"
            ingested_txt = 0
            ingested_json = 0
            try:
                if ingested_dir.exists() and ingested_dir.is_dir():
                    ingested_txt = len(list(ingested_dir.glob("*.txt")))
                    ingested_json = len(list(ingested_dir.glob("*.json")))
            except Exception:
                # Non bloquant (droits/FS)
                pass

            self.source_stats = {
                "source_dir": str(input_dir),
                "total_files": len(files),
                "by_ext": dict(sorted(ext_counts.items(), key=lambda kv: (-kv[1], kv[0]))),
                "audio_ingested": {"txt": ingested_txt, "json": ingested_json},
            }

            # Log utilisateur (simple) + push vers le frontend
            pdf_n = ext_counts.get(".pdf", 0)
            docx_n = ext_counts.get(".docx", 0)
            txt_n = ext_counts.get(".txt", 0)
            m4a_n = ext_counts.get(".m4a", 0)
            mp3_n = ext_counts.get(".mp3", 0)
            wav_n = ext_counts.get(".wav", 0)
            self._log_progress(
                "EXTRACTING",
                (
                    f"Sources détectées: {len(files)} fichier(s) "
                    f"(pdf={pdf_n}, docx={docx_n}, txt={txt_n}, m4a={m4a_n}, mp3={mp3_n}, wav={wav_n}). "
                    f"Audio RAG (transcriptions): {ingested_txt} .txt"
                ),
                0.12,
            )

            # Extraction de chaque fichier
            docs: List[ExtractedDoc] = []
            for file_path in files:
                doc = extract_one(file_path, self.params.enable_soffice)
                if doc.text:  # On ne garde que les docs avec du texte
                    docs.append(doc)

            # Stats post-extraction
            if self.source_stats is None:
                self.source_stats = {}
            self.source_stats["extracted_docs"] = len(docs)

            if not docs:
                raise ValueError(f"Aucun document extrait avec succès dans {input_dir}")
            return docs

        # Déterminer le dossier source
        fallback_dir: Optional[Path] = None
        if self.params.source_file:
            source_path = Path(self.params.source_file)
            if not source_path.exists():
                raise FileNotFoundError(f"Fichier source introuvable: {source_path}")
            input_dir = source_path.parent
        else:
            # Par défaut, chercher un dossier "sources" (si c'est une vraie arborescence de sources)
            # mais si ce dossier est vide/inutile, retomber sur le dossier client.
            sources_dir = self.params.client_dir / "sources"
            if sources_dir.exists():
                input_dir = sources_dir
                fallback_dir = self.params.client_dir
            else:
                input_dir = self.params.client_dir

        # Optionnel: ingestion audio automatique avant de lister / extraire.
        # On le fait une seule fois (sinon, avec le fallback sources->client_dir,
        # on pourrait rescanner 2x le même dossier).
        try:
            audio_summary = _maybe_auto_ingest_audio()
            if audio_summary.get("processed"):
                self._log_progress(
                    "EXTRACTING",
                    f"🎙️ Audio: {audio_summary.get('processed')} transcription(s) générée(s) automatiquement.",
                    0.11,
                )
        except Exception:
            # Non bloquant
            pass

        # Tentative 1 (dossier choisi)
        try:
            docs = _attempt_extract(input_dir)
        except ValueError as exc:
            if fallback_dir and fallback_dir != input_dir:
                logger.info("Extraction vide dans %s, fallback sur %s (%s)", input_dir, fallback_dir, exc)
                self._log_progress(
                    "EXTRACTING",
                    f"Aucun document utilisable dans {input_dir}. Nouvelle tentative dans {fallback_dir}…",
                    0.12,
                )
                docs = _attempt_extract(fallback_dir)
                input_dir = fallback_dir
            else:
                raise
        
        # Sauvegarde du JSON
        payload = {
            "documents": [
                {
                    "path": doc.path,
                    "ext": doc.ext,
                    "hash": doc.text_sha256,
                    "text": doc.text,
                    "pages": doc.pages,
                    "mtime": doc.mtime_iso,
                    "extractor": doc.extractor,
                    "size_bytes": doc.size_bytes
                }
                for doc in docs
            ],
            "metadata": {
                "client": self.params.client_dir.name,
                "source_dir": str(input_dir),
                "doc_count": len(docs),
                "source_stats": self.source_stats,
            }
        }
        
        extracted_path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8"
        )
        
        return extracted_path
    
    def _generate_fields(self, extracted_path: Path) -> Path:
        """Génération des champs via LLM."""
        answers_path = self.temp_dir / "answers.json"
        
        # Charger le payload extrait
        payload = json.loads(extracted_path.read_text(encoding="utf-8"))
        
        # Extraire les placeholders du template
        placeholders = extract_placeholders_from_docx(self.params.template_path)
        if not placeholders:
            raise ValueError("Aucun placeholder {{...}} trouvé dans le template")
        
        # Construire les field specs
        fields = build_field_specs(placeholders, fallback_defs=None)

        # Initialiser la progression par champ (pour affichage côté UI)
        self._init_field_progress(fields)
        if self.progress_callback:
            self.progress_callback(
                {
                    "status": "GENERATING",
                    "message": "Initialisation de la progression des champs…",
                    "progress": 0.4,
                    "field_progress": self.field_progress,
                    "field_order": self.field_order,
                    "field_progress_version": self.field_progress_version,
                }
            )
        
        # Détecter le numéro AVS si non fourni
        avs_value = self.params.avs_number
        if not avs_value:
            detected_avs = detect_avs_number(payload)
            if detected_avs:
                avs_value = detected_avs
                logger.info(f"Numéro AVS détecté: {avs_value}")
        
        # Construire la valeur LIEU_ET_DATE
        location_date_value = build_location_date(
            self.params.location_city,
            self.params.location_date,
            auto_date=self.params.auto_location_date,
            date_format=self.params.date_format
        )
        
        # Champs déterministes
        deterministic = {
            "MONSIEUR_OU_MADAME": self.params.civility,
            "NAME": self.params.name,
            "SURNAME": self.params.surname,
            "LIEU_ET_DATE": location_date_value,
            "NUMERO_AVS": avs_value,
        }
        
        # Génération des champs
        answers = generate_fields(
            payload,
            model=self.params.llm_model,
            host=self.params.llm_host,
            temperature=self.params.temperature,
            topk=self.params.topk,
            top_p=self.params.top_p,
            fields=fields,
            deterministic_values=deterministic,
            include_filters=self.params.include_filters or [],
            exclude_filters=self.params.exclude_filters or [],
            debug_dir=self.temp_dir / "debug",
            status_callback=lambda msg: self._log_progress("GENERATING", msg, None),
            progress_callback=lambda k, stg, m: self._update_field_progress(k, stg, m),
        )
        
        # Sauvegarde
        answers_path.write_text(
            json.dumps(answers, ensure_ascii=False, indent=2),
            encoding="utf-8"
        )
        
        return answers_path
    
    def _render_document(self, answers_path: Path):
        """Rendu du document Word."""
        # Créer le dossier de sortie si nécessaire
        self.params.output_path.parent.mkdir(parents=True, exist_ok=True)
        
        # Import locaux pour éviter les problèmes circulaires
        import json
        from docx import Document
        from core.render import replace_text_everywhere, build_moustache_mapping
        
        # Charger les réponses
        answers_dict = json.loads(answers_path.read_text(encoding="utf-8"))
        
        # Ouvrir le template
        doc = Document(str(self.params.template_path))
        
        # Remplacement simple des champs déterministes
        simple_mapping = {
            "{{MONSIEUR_OU_MADAME}}": self.params.civility,
            "{{NAME}}": self.params.name,
            "{{SURNAME}}": self.params.surname,
        }
        replace_text_everywhere(doc, simple_mapping)
        
        # Remplacement des champs générés par le LLM
        moustache_mapping = build_moustache_mapping(answers_dict)
        if moustache_mapping:
            replace_text_everywhere(doc, moustache_mapping)
        
        # Sauvegarder
        doc.save(str(self.params.output_path))
        logger.info(f"Document rendu: {self.params.output_path}")
    
    def _export_pdf(self):
        """Export PDF via LibreOffice (optionnel)."""
        import subprocess
        
        pdf_path = self.params.output_path.with_suffix(".pdf")
        
        try:
            subprocess.run([
                "soffice",
                "--headless",
                "--convert-to", "pdf",
                "--outdir", str(pdf_path.parent),
                str(self.params.output_path)
            ], check=True, capture_output=True, timeout=30)
            
            logger.info(f"PDF exporté: {pdf_path}")
        except Exception as e:
            logger.warning(f"Échec de l'export PDF: {e}")
